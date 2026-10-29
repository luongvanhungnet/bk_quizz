"""Meaningful geometry, lineage and worker contract regressions; no model downloads."""
import json
import importlib.util
import math
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from adapters import (normalized_box,import_docling,import_mineru,import_paddle,
                      mineru_command,run_docling,run_mineru,file_hash,extract_pptx)
from paddle_worker import infer_task,result_dict,safe_path,task_geometry,run


class Adapters(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def write(self,name,value):
        path=self.root/name
        path.write_text(json.dumps(value,ensure_ascii=False),encoding='utf-8')
        return path

    def test_top_left_geometry(self):
        self.assertEqual(normalized_box([10,20,50,80],100,200),[100,100,500,400])

    def test_bottom_left_docling_geometry(self):
        b={'l':10,'t':180,'r':50,'b':120,'coord_origin':'BOTTOMLEFT'}
        self.assertEqual(normalized_box(b,100,200),[100,100,500,400])

    def test_invalid_geometry_not_clamped(self):
        for box in ([0,0,101,100],[True,0,50,50],[0,0,0,50],[0,0,math.nan,20],[-1,0,50,50]):
            self.assertIsNone(normalized_box(box,100,100))

    def test_no_guess_of_units(self):
        self.assertIsNone(normalized_box([0,0,1,1],units='unknown'))
        self.assertEqual(normalized_box([0,0,1,1],units='normalized_1'),[0,0,1000,1000])

    def docling_fixture(self):
        return {'metadata':{'source_sha256':'f'*64,'do_ocr':False,'version':'fixture'},'document':{
            'pages':{'1':{'size':{'width':100,'height':200}}},
            'body':{'children':[{'$ref':'#/texts/1'},{'$ref':'#/texts/0'}]},
            'texts':[{'self_ref':'#/texts/0','label':'formula','text':'x ≤ y',
                'prov':[{'page_no':1,'bbox':{'l':10,'t':180,'r':50,'b':120,'coord_origin':'BOTTOMLEFT'}}]},
                {'self_ref':'#/texts/1','label':'text','text':'Thuật toán ứng dụng',
                'prov':[{'page_no':1,'bbox':{'l':0,'t':0,'r':100,'b':10,'coord_origin':'TOPLEFT'}}]}]}}

    def test_docling_body_reading_order_not_storage_order(self):
        path=self.write('docling.json',self.docling_fixture())
        rows=import_docling(path)
        self.assertEqual([r['text'] for r in rows],['Thuật toán ứng dụng','x ≤ y'])
        self.assertEqual(rows[1]['bbox'],[100,100,500,400])
        self.assertEqual(rows[0]['evidence']['backend_family'],'docling_native_layout')
        self.assertFalse(rows[0]['evidence']['semantic_verified'])
        self.assertEqual(rows[0]['evidence']['source_sha256'],'f'*64)

    def test_docling_cyclic_refs_terminate(self):
        payload=self.docling_fixture()
        payload['document']['texts'][0]['children']=[{'$ref':'#/body'}]
        self.assertEqual(len(import_docling(self.write('cyclic.json',payload))),2)

    def test_docling_multiple_page_provenance_slices_charspan(self):
        payload=self.docling_fixture()
        payload['document']['body']['children']=[]
        item=payload['document']['texts'][0]
        item['text']='ABCDEF';item['prov'][0]['charspan']=[0,3]
        item['prov'].append({'page_no':2,'bbox':{'l':0,'t':0,'r':20,'b':20,'coord_origin':'TOPLEFT'},'charspan':[3,6]})
        payload['document']['pages']['2']={'size':{'width':100,'height':200}}
        rows=import_docling(self.write('multi.json',payload))
        self.assertEqual([r['text'] for r in rows[:2]],['ABC','DEF'])
        self.assertEqual([r['page'] for r in rows[:2]],[1,2])

    def test_docling_unknown_page_geometry_kept_unknown(self):
        payload=self.docling_fixture()
        payload['document']['pages']={}
        rows=import_docling(self.write('unknown.json',payload))
        self.assertTrue(all(r['bbox'] is None for r in rows))

    def test_docling_table_cell_text_escaped_and_spans_preserved(self):
        payload=self.docling_fixture();doc=payload['document']
        doc['tables']=[{'label':'table','data':{'table_cells':[{'start_row_offset_idx':0,'start_col_offset_idx':0,
            'text':'a < 3','row_span':2,'col_span':1}]},'prov':doc['texts'][0]['prov']}]
        row=import_docling(self.write('table.json',payload))[-1]
        self.assertIn('rowspan="2"',row['text']);self.assertIn('a &lt; 3',row['text'])

    def test_mineru_zero_based_pages_and_math_symbols(self):
        path=self.write('content_list.json',[{'type':'text','page_idx':0,'bbox':[1,2,99,100],'text':'x ≤ y; O(n²)'}])
        row=import_mineru(path)[0]
        self.assertEqual(row['page'],1);self.assertEqual(row['bbox'],[1,2,99,100])
        self.assertEqual(row['text'],'x ≤ y; O(n²)')

    def test_mineru_missing_page_is_not_guessed(self):
        row=import_mineru(self.write('no-page.json',[{'type':'text','text':'abc','bbox':[0,0,100,100]}]))[0]
        self.assertIsNone(row['page']);self.assertEqual(row['status'],'error')

    def test_mineru_backend_lineage(self):
        row=import_mineru(self.write('miner.json',{'metadata':{'backend':'pipeline'},'content_list':[
            {'type':'table','page_idx':2,'bbox':[0,0,20,20],'table_body':'<table><tr><td>2</td></tr></table>'}]}))[0]
        self.assertEqual(row['page'],3);self.assertEqual(row['kind'],'table')
        self.assertEqual(row['evidence']['backend_family'],'mineru_pipeline_paddle')

    def test_mineru_modern_units_explicit_only(self):
        payload={'blocks':[{'page':1,'bbox':[0,0,1,1],'text':'x'}]}
        self.assertIsNone(import_mineru(self.write('modern1.json',payload))[0]['bbox'])
        payload['metadata']={'coord_system':'normalized_1','page_base':1}
        self.assertEqual(import_mineru(self.write('modern2.json',payload))[0]['bbox'],[0,0,1000,1000])

    def test_mineru_unsupported_markdown_not_silently_accepted(self):
        with self.assertRaises(ValueError):
            import_mineru(self.write('bad.json',{'markdown':'# header'}))

    def modern_fixture(self):
        return {'schema':'docvortex.middle','schema_version':'2.0',
            'metadata':{'file_suffix':'pdf','producer':{'name':'mineru','version':'4.0.12'}},
            'extensions':{'mineru':{'tier':'standard','parse_mode':'ocr'},
                'docvortex_layout':{'version':1,'pages':[{'page_idx':3,'width_pt':600,'height_pt':400}]}},
            'pages':[{'page_idx':3,'blocks':[{'type':'text','index':0,'bbox':[60,40,300,120],
                'content':[{'type':'text','content':'Độ phức tạp '},{'type':'equation_inline','content':'O(n^2)'}]},
                {'type':'table','bbox':[60,160,540,360],'content':[{'type':'table_body','content':'<table><tr><td>12</td></tr></table>'},
                    {'type':'table_caption','bbox':[60,360,540,380],'content':[{'type':'text','content':'Bảng 1'}]}]}]}]}

    def test_docvortex_current_schema_points_and_source_page(self):
        payload={'metadata':{'source_sha256':'f'*64,'source_file':'original.pdf'},'document':self.modern_fixture()}
        rows=import_mineru(self.write('middle_json.json',payload))
        self.assertEqual(rows[0]['page'],4);self.assertEqual(rows[0]['bbox'],[100,100,500,300])
        self.assertEqual(rows[0]['text'],'Độ phức tạp O(n^2)')
        self.assertEqual(rows[0]['evidence']['source_sha256'],'f'*64)
        self.assertEqual(rows[0]['evidence']['parser_version'],'4.0.12')
        self.assertEqual(rows[0]['evidence']['backend_family'],'mineru_standard_ocr')
        self.assertEqual(rows[1]['kind'],'table');self.assertEqual(rows[2]['text'],'Bảng 1')
        self.assertNotIn('Bảng 1',rows[1]['text'])

    def test_docvortex_missing_layout_keeps_unknown_geometry(self):
        payload=self.modern_fixture();payload['extensions'].pop('docvortex_layout')
        self.assertTrue(all(r['bbox'] is None for r in import_mineru(self.write('middle_missing.json',payload))))

    def test_docvortex_unknown_protocol_rejected(self):
        payload=self.modern_fixture();payload['schema_version']='9.9'
        with self.assertRaises(ValueError):
            import_mineru(self.write('future.json',payload))

    def test_legacy_mineru_cli_no_shell(self):
        cmd=mineru_command('C:/Program Files/mineru.exe','a.pdf','b','--path --output --backend')
        self.assertEqual(cmd,['C:/Program Files/mineru.exe','-p','a.pdf','-o','b','-b','pipeline'])

    def test_modern_mineru_tier_not_misrepresented_as_legacy(self):
        with self.assertRaises(ValueError):
            mineru_command('mineru-kit','a.pdf','b','mineru-kit parse')
        cmd=mineru_command('mineru-kit','a.pdf','b','mineru-kit parse',backend='standard')
        self.assertIn('standard',cmd);self.assertIn('--format',cmd)

    def test_unknown_cli_is_rejected(self):
        with self.assertRaises(ValueError):
            mineru_command('mineru','a.pdf','b','Usage: mineru find search')

    def test_paddle_structure_crop_geometry_maps_to_full_slide(self):
        payload={'metadata':{'page':5,'image_width':200,'image_height':100,'crop_bbox':[200,100,600,500]},
            'res':{'parsing_res_list':[{'block_bbox':[0,0,100,50],'block_label':'text','block_content':'Đồ thị G=(V,E)'}]}}
        row=import_paddle(self.write('paddle.json',payload))[0]
        self.assertEqual(row['bbox'],[200,100,400,300]);self.assertEqual(row['page'],5)

    def test_paddle_image_crop_not_assumed_full_page(self):
        payload={'metadata':{'page':1,'image_width':200,'image_height':100},
            'res':{'parsing_res_list':[{'block_bbox':[0,0,100,50],'block_label':'text','block_content':'abc'}]}}
        self.assertIsNone(import_paddle(self.write('unframed.json',payload))[0]['bbox'])

    def test_paddle_page_index_zero_based_and_full_frame(self):
        row=import_paddle(self.write('full.json',{'metadata':{'image_width':100,'image_height':100,'is_full_page':True},
            'res':{'page_index':0,'rec_texts':['x²'],'rec_boxes':[[0,0,50,50]],'rec_scores':[.8]}}))[0]
        self.assertEqual(row['page'],1);self.assertEqual(row['bbox'],[0,0,500,500]);self.assertEqual(row['confidence'],.8)

    def test_paddle_no_duplicate_ocr_when_structure_has_text(self):
        path=self.write('dup.json',{'metadata':{'page':1},'res':{'parsing_res_list':[
            {'block_label':'text','block_content':'abc'}],'rec_texts':['abc']}})
        self.assertEqual(len(import_paddle(path)),1)

    def test_paddle_standalone_formula_and_table(self):
        payload={'metadata':{'page':1,'crop_bbox':[100,100,200,200],'image_width':10,'image_height':10},
            'records':[{'metadata':{'page':1,'crop_bbox':[100,100,200,200]},'res':{'rec_formula':'x^2'}},
                {'metadata':{'page':1,'image_width':10,'image_height':10,'is_full_page':True},
                 'res':{'table_res_list':[{'bbox':[0,0,10,10],'pred_html':'<table/>'}]}}]}
        rows=import_paddle(self.write('specialist.json',payload))
        self.assertEqual([r['kind'] for r in rows],['formula','table'])
        self.assertEqual(rows[0]['bbox'],[100,100,200,200])

    def test_docling_unavailable_is_durable_not_fake_success(self):
        pdf=self.root/'a.pdf';pdf.write_bytes(b'%PDF-canary')
        with patch.dict('sys.modules',{'docling':None}):
            result=run_docling(pdf,self.root/'unavailable.json')
        self.assertEqual(result['status'],'unavailable')
        self.assertIn('traceback',result);self.assertEqual(result['metadata']['source_sha256'],file_hash(pdf))

    def test_mineru_unavailable_is_durable(self):
        pdf=self.root/'a.pdf';pdf.write_bytes(b'%PDF')
        with patch('adapters.subprocess.run',side_effect=FileNotFoundError('not installed')):
            result=run_mineru(pdf,self.root/'mineru')
        self.assertEqual(result['status'],'unavailable');self.assertTrue((self.root/'mineru'/'run.json').exists())

    @unittest.skipUnless(importlib.util.find_spec('pptx'),'Optional python-pptx not installed')
    def test_native_pptx_real_file_canary_text_merge_chart(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(.5));shape.text='Thuật toán: x ≤ 3'
        table=slide.shapes.add_table(2,2,Inches(1),Inches(2),Inches(4),Inches(1)).table
        table.cell(0,0).text='12';table.cell(0,0).merge(table.cell(1,0));table.cell(0,1).text='7';table.cell(1,1).text='8'
        data=CategoryChartData();data.categories=['a','b'];data.add_series('Số bước',[1,2])
        slide.shapes.add_chart(XL_CHART_TYPE.LINE,Inches(1),Inches(3),Inches(4),Inches(2),data)
        path=self.root/'original.pptx';deck.save(path)
        rows=extract_pptx(path)
        self.assertEqual([r['kind'] for r in rows],['text','table','chart'])
        self.assertEqual(rows[0]['text'],'Thuật toán: x ≤ 3')
        self.assertIn('rowspan="2"',rows[1]['text']);self.assertEqual(rows[1]['text'].count('12'),1)
        self.assertEqual(json.loads(rows[2]['text'])['series'][0]['values'],[1,2])
        self.assertTrue(all(r['evidence']['backend_family']=='native_office' for r in rows))


class Worker(unittest.TestCase):
    def setUp(self):
        from PIL import Image
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        Image.new('RGB',(100,50),'white').save(self.root/'crop.png')
        self.task={'id':'p1_formula','page':1,'image':'crop.png','crop_bbox':[100,100,600,400],
                   'kind':'formula','native_text':'x² ≤ y','source_sha256':'f'*64}

    def tearDown(self):
        self.temp.cleanup()

    def pipeline(self,raw):
        class Fake:
            def predict(self,path):
                return [raw]
        return Fake()

    def test_formula_only_explicit_kind(self):
        task={**self.task,'kind':'text'}
        result=infer_task(task,self.root,self.root/'out',mode='formula',pipeline=self.pipeline({'rec_formula':'x^2'}))
        self.assertEqual(result['status'],'error');self.assertIn('kind=formula',result['error'])

    def test_formula_preserves_raw_native_and_symbols(self):
        result=infer_task(self.task,self.root,self.root/'out',mode='formula',pipeline=self.pipeline({'rec_formula':r'x^2 \le y'}))
        self.assertEqual(result['status'],'ok');self.assertEqual(result['native_text'],'x² ≤ y')
        self.assertEqual(result['candidates'][0]['bbox'],[100,100,600,400])
        self.assertFalse(result['semantic_verified'])

    def test_original_colon_id_stays_in_evidence_with_safe_filename(self):
        task={**self.task,'id':'c6c829f14bf2:p0020:text'}
        result=infer_task(task,self.root,self.root/'out',mode='formula',pipeline=self.pipeline({'rec_formula':'x'}))
        self.assertEqual(result['candidates'][0]['evidence']['region_id'],task['id'])
        self.assertTrue(all(':' not in p.name for p in (self.root/'out'/'records').iterdir()))

    def test_formula_empty_fails_with_traceback(self):
        result=infer_task(self.task,self.root,self.root/'out',mode='formula',pipeline=self.pipeline({'rec_formula':''}))
        self.assertEqual(result['status'],'error');self.assertIn('ValueError',result['traceback'])

    def test_worker_cache_does_not_rerun(self):
        out=self.root/'out'
        a=infer_task(self.task,self.root,out,mode='formula',pipeline=self.pipeline({'rec_formula':'x'}))
        with patch('paddle_worker.create_pipeline',side_effect=AssertionError('Must not load')):
            b=infer_task(self.task,self.root,out,mode='formula')
        self.assertEqual(a,b)

    def test_changed_image_cannot_reuse_cached_result(self):
        out=self.root/'out'
        infer_task(self.task,self.root,out,mode='formula',pipeline=self.pipeline({'rec_formula':'x'}))
        (self.root/'crop.png').write_bytes(b'changed')
        with self.assertRaises(FileExistsError):
            infer_task(self.task,self.root,out,mode='formula')

    def test_expected_hash_checked(self):
        with self.assertRaisesRegex(ValueError,'hash mismatch'):
            infer_task({**self.task,'input_file_sha256':'0'*64},self.root,self.root/'out',mode='formula')

    def test_path_traversal_rejected(self):
        with self.assertRaises(ValueError):
            safe_path(self.root,'../secret.png')

    def test_task_geometry_not_guessed(self):
        with self.assertRaises(ValueError):
            task_geometry({'page':1},(100,50))
        self.assertEqual(task_geometry({'page':1,'is_full_page':True},(100,50))['crop_bbox'],[0,0,1000,1000])

    def test_table_does_not_fabricate_table_from_ocr(self):
        task={**self.task,'kind':'table'}
        result=infer_task(task,self.root,self.root/'out',mode='table',pipeline=self.pipeline({'rec_texts':['abc']}))
        self.assertEqual(result['status'],'error');self.assertIn('no nonempty table',result['error'])

    def test_json_string_wrapped_result(self):
        class Result:
            json='{"res":{"rec_formula":"x^2"}}'
        self.assertEqual(result_dict(Result()),{'rec_formula':'x^2'})

    def test_worker_initialization_failure_only_attempted_once(self):
        task2={**self.task,'id':'second'}
        plan=self.root/'plan.json';plan.write_text(json.dumps({'requests':[self.task,task2]}),encoding='utf-8')
        expected={'paddleocr':'3.7.0','paddlex':'3.7.2','paddlepaddle':'3.2.2','paddlepaddle-gpu':None}
        with patch('paddle_worker.package_version',side_effect=lambda name:expected[name]),\
             patch('paddle_worker.create_pipeline',side_effect=RuntimeError('Model missing')) as create:
            result=run(plan,self.root/'out',mode='formula')
        self.assertEqual(create.call_count,1);self.assertEqual(len(result['records']),2)
        self.assertEqual(result['status'],'error')
        self.assertIn('previously failed',result['records'][1]['error'])

    def test_valid_specialist_seed_skips_model_and_rerun_is_explicit(self):
        plan=self.root/'plan.json'
        plan.write_text(json.dumps({'requests':[{**self.task,'specialist_seed_valid':True}]}),encoding='utf-8')
        expected={'paddleocr':'3.7.0','paddlex':'3.7.2','paddlepaddle':'3.2.2','paddlepaddle-gpu':None}
        with patch('paddle_worker.package_version',side_effect=lambda name:expected[name]),\
             patch('paddle_worker.create_pipeline',side_effect=AssertionError('Must skip')) as create:
            result=run(plan,self.root/'out',mode='formula')
        self.assertEqual(result['request_count'],0);self.assertFalse(create.called)

    def test_ocr_only_routes_text_or_explicit_request(self):
        plan=self.root/'plan.json'
        plan.write_text(json.dumps({'requests':[self.task,{**self.task,'id':'text','kind':'text'}]}),encoding='utf-8')
        expected={'paddleocr':'3.7.0','paddlex':'3.7.2','paddlepaddle':'3.2.2','paddlepaddle-gpu':None}
        with patch('paddle_worker.package_version',side_effect=lambda name:expected[name]),\
             patch('paddle_worker.create_pipeline',return_value=self.pipeline({'rec_texts':['x'],'rec_boxes':[[0,0,100,50]]})):
            result=run(plan,self.root/'out',mode='ocr')
        self.assertEqual(result['request_count'],1);self.assertEqual(result['records'][0]['id'],'text')


if __name__=='__main__':
    unittest.main()
