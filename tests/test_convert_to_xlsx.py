"""Round-trip checks for the command-line data converter."""
from pathlib import Path
import json
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import convert_to_xlsx as converter
from openpyxl import load_workbook


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.output = self.root / "output.xlsx"

    def jsonl(self, rows):
        p = self.root / "input.jsonl"
        p.write_text("\n".join(json.dumps(row, ensure_ascii=False) for row in rows), encoding="utf-8")
        return p

    def read(self, path=None):
        book = load_workbook(path or self.output)
        self.addCleanup(book.close)
        return book

    def test_nested_arrays_new_fields_and_ambiguous_keys(self):
        source = self.jsonl([
            {"id": "001", "a.b": "literal", "a": {"b": "nested"}, "options": [{"text": "Có", "correct": True}], "tokens": [1, 2]},
            {"id": "002", "new_field": "xuất hiện cuối", "empty": {}},
        ])
        converter.convert(source, self.output)
        sheet = self.read().active
        rows = list(sheet.values)
        first, second = [dict(zip(rows[0], row)) for row in rows[1:]]
        self.assertEqual(first['["a.b"]'], "literal")
        self.assertEqual(first["a.b"], "nested")
        self.assertEqual(first["options[0].text"], "Có")
        self.assertIs(first["options[0].correct"], True)
        self.assertEqual(first["tokens"], "[1, 2]")
        self.assertEqual(second["new_field"], "xuất hiện cuối")
        self.assertEqual(second["empty"], "{}")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertTrue(sheet.auto_filter.ref.endswith("3"))

    def test_long_unicode_text_is_preserved_across_columns(self):
        text = "Tiếng Việt 🙂\n" * 5_000
        converter.convert(self.jsonl([{"text": text}]), self.output)
        values = list(self.read().active.values)[1]
        self.assertGreater(len(values), 1)
        self.assertEqual("".join(values), text)
        self.assertTrue(all(len(v.encode("utf-16-le")) // 2 <= 32767 for v in values))

    def test_csv_bom_delimiter_multiline_and_literal_formula(self):
        source = self.root / "input.csv"
        source.write_text('id;note;expression\n00123;"dòng 1\ndòng 2";=1+1\n', encoding="utf-8-sig")
        converter.convert(source, self.output)
        sheet = self.read().active
        self.assertEqual(sheet["A2"].value, "00123")
        self.assertEqual(sheet["B2"].value, "dòng 1\ndòng 2")
        self.assertEqual(sheet["C2"].value, "=1+1")
        self.assertEqual(sheet["C2"].data_type, "s")

    def test_csv_duplicate_empty_headers_keep_all_values(self):
        source = self.root / "input.csv"
        source.write_text(',id,id\nx,001,002\n', encoding="utf-8")
        converter.convert(source, self.output, delimiter=",")
        rows = list(self.read().active.values)
        self.assertEqual(rows[0], ("column_1", "id", "id (2)"))
        self.assertEqual(rows[1], ("x", "001", "002"))

    def test_windows_line_endings_inside_json_cells(self):
        converter.convert(self.jsonl([{"text": "dòng 1\r\ndòng 2\rdòng 3\ndòng 4"}]), self.output)
        self.assertEqual(self.read().active["A2"].value, "dòng 1\ndòng 2\ndòng 3\ndòng 4")

    def test_csv_explicit_encoding_and_tab(self):
        source = self.root / "input.csv"
        source.write_text('name\tvalue\nTiếng Việt\t001\n', encoding="utf-16")
        converter.convert(source, self.output, encoding="utf-16", delimiter="\t")
        self.assertEqual(list(self.read().active.values)[1], ("Tiếng Việt", "001"))

    def test_large_numbers_and_types(self):
        source = self.root / "input.jsonl"
        source.write_text('{"big":1234567890123456789,"precise":0.1234567890123456789,"small":0.1,"tiny":1e-500,"number":123,"boolean":true,"nothing":null,"formula":"=1+1"}', encoding="utf-8")
        converter.convert(source, self.output)
        sheet = self.read().active
        self.assertEqual(sheet["A2"].value, "1234567890123456789")
        self.assertEqual(sheet["B2"].value, "0.1234567890123456789")
        self.assertEqual(sheet["C2"].value, 0.1)
        self.assertEqual(sheet["D2"].value, "1E-500")
        self.assertEqual(sheet["E2"].value, 123)
        self.assertEqual(sheet["F2"].data_type, "b")
        self.assertIsNone(sheet["G2"].value)
        self.assertEqual(sheet["H2"].data_type, "s")

    def test_default_output_uses_terminal_directory(self):
        source_dir, cwd = self.root / "nguồn dữ liệu", self.root / "terminal folder"
        source_dir.mkdir()
        cwd.mkdir()
        source = source_dir / "train.samples.jsonl"
        original = '{"text":"Câu hỏi tiếng Việt"}\n'.encode("utf-8")
        source.write_bytes(original)
        script = Path(converter.__file__).resolve()
        run = subprocess.run([sys.executable, str(script), str(source)], cwd=cwd, capture_output=True, encoding="utf-8")
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertTrue((cwd / "train.samples.xlsx").is_file())
        self.assertFalse((source_dir / "train.samples.xlsx").exists())
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(self.read(cwd / "train.samples.xlsx").active["A2"].value, "Câu hỏi tiếng Việt")

    def test_overwrite_requires_force(self):
        source = self.jsonl([{"value": "new"}])
        self.output.write_bytes(b"keep original")
        with self.assertRaises(FileExistsError):
            converter.convert(source, self.output)
        self.assertEqual(self.output.read_bytes(), b"keep original")
        converter.convert(source, self.output, force=True)
        self.assertEqual(self.read().active["A2"].value, "new")

    def test_invalid_json_does_not_publish_partial_workbook(self):
        for bad in ['{"x":', '{"x":1,"x":2}', '{"x":NaN}']:
            with self.subTest(bad=bad):
                source = self.root / "input.jsonl"
                source.write_text('{"ok":1}\n' + bad, encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "dòng 2"):
                    converter.convert(source, self.output)
                self.assertFalse(self.output.exists())

    def test_invalid_csv_keeps_existing_output_even_with_force(self):
        source = self.root / "input.csv"
        source.write_text('a,b\n1,2,3\n', encoding="utf-8")
        self.output.write_bytes(b"keep original")
        with self.assertRaisesRegex(ValueError, "CSV dòng 2"):
            converter.convert(source, self.output, force=True, delimiter=",")
        self.assertEqual(self.output.read_bytes(), b"keep original")

    def test_excel_limits_and_invalid_xml(self):
        with patch.object(converter, "MAX_DATA_ROWS", 2):
            converter.convert(self.jsonl([{"id": i} for i in range(5)]), self.output)
        book = self.read()
        self.assertEqual(book.sheetnames, ["Data", "Data_2", "Data_3"])
        self.assertEqual([row[0] for ws in book for row in list(ws.values)[1:]], list(range(5)))
        other = self.root / "invalid.xlsx"
        with self.assertRaises(ValueError):
            converter.convert(self.jsonl([{"bad": "a\0b"}]), other)
        self.assertFalse(other.exists())
        with patch.object(converter, "MAX_COLUMNS", 1), self.assertRaises(ValueError):
            converter.convert(self.jsonl([{"a": 1, "b": 2}]), other)

    def test_empty_input_reports_error(self):
        source = self.root / "empty.jsonl"
        source.write_text('\n\n', encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "không có bản ghi"):
            converter.convert(source, self.output)
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
