#!/usr/bin/env python3
"""Convert JSONL/NDJSON or CSV to XLSX in the terminal's current directory.

Requires Python >= 3.10 and openpyxl (see requirements-xlsx.txt).
"""
from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from decimal import Decimal
import json
import math
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile

MAX_CELL_UNITS = 32_767
MAX_COLUMNS = 16_384
MAX_DATA_ROWS = 1_048_575  # Excel's row limit minus the header.
INVALID_XML = re.compile(r"[^\x09\x0A\x0D\x20-\uD7FF\uE000-\uFFFD\U00010000-\U0010FFFF]")


def key_path(parent: str, key: str) -> str:
    if key.isidentifier():
        return f"{parent}.{key}" if parent else key
    return parent + "[" + json.dumps(key, ensure_ascii=False) + "]"


def json_text(value) -> str:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, list):
        return "[" + ", ".join(json_text(item) for item in value) + "]"
    if isinstance(value, dict):
        return "{" + ", ".join(json.dumps(k, ensure_ascii=False) + ": " + json_text(v) for k, v in value.items()) + "}"
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def flatten(value, parent: str = "") -> dict:
    """Expand objects/object arrays; keep scalar arrays as readable JSON text."""
    if isinstance(value, dict) and value:
        result = {}
        for key, child in value.items():
            result.update(flatten(child, key_path(parent, key)))
        return result
    if isinstance(value, list) and value and all(isinstance(x, dict) for x in value):
        result = {}
        for index, child in enumerate(value):
            result.update(flatten(child, f"{parent or 'value'}[{index}]"))
        return result
    if isinstance(value, (list, dict)):
        value = json_text(value)
    return {parent or "value": value}


def unique_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Trùng khóa JSON: {key!r}")
        result[key] = value
    return result


def invalid_constant(value):
    raise ValueError(f"Giá trị không hợp lệ trong JSON: {value}")


def unique_name(name: str, used: set[str]) -> str:
    candidate, number = name, 2
    while candidate in used:
        candidate = f"{name} ({number})"
        number += 1
    used.add(candidate)
    return candidate


def records(path: Path, encoding: str, delimiter: str | None):
    with path.open("r", encoding=encoding, newline="") as stream:
        if path.suffix.lower() in {".jsonl", ".ndjson"}:
            for line_number, line in enumerate(stream, 1):
                if not line.strip():
                    continue
                try:
                    obj = json.loads(line, parse_float=Decimal, object_pairs_hook=unique_object,
                                     parse_constant=invalid_constant)
                    yield line_number, flatten(obj)
                except (ValueError, RecursionError) as exc:
                    raise ValueError(f"JSONL dòng {line_number}: {exc}") from exc
        else:
            csv.field_size_limit(2_147_483_647)
            sample = stream.read(65_536)
            stream.seek(0)
            if delimiter is None:
                try:
                    delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                except csv.Error:
                    delimiter = ","
            reader = csv.reader(stream, delimiter=delimiter, strict=True)
            try:
                header = next(reader, None)
                while header == []:
                    header = next(reader, None)
                if header is None:
                    return
                used = set()
                header = [unique_name(name or f"column_{i}", used) for i, name in enumerate(header, 1)]
                for row in reader:
                    if not row:
                        continue
                    if len(row) != len(header):
                        raise ValueError(f"CSV dòng {reader.line_num}: có {len(row)} ô, cần {len(header)}. Kiểm tra --delimiter.")
                    yield reader.line_num, dict(zip(header, row))
            except csv.Error as exc:
                raise ValueError(f"CSV dòng {reader.line_num}: {exc}") from exc


def excel_value(value):
    if isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > 15:
        return str(value)
    if isinstance(value, Decimal):
        number = float(value)
        if len(value.as_tuple().digits) > 15 or not math.isfinite(number) or (value and abs(number) < sys.float_info.min):
            return str(value)
        return number
    return value


def cell_parts(value, label: str) -> list:
    value = excel_value(value)
    if not isinstance(value, str):
        return [value]
    # Excel uses LF inside cells. Normalize CRLF/CR before the XML writer;
    # otherwise Windows text streams can introduce doubled line breaks.
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    if INVALID_XML.search(value):
        raise ValueError(f"{label}: chứa ký tự điều khiển hoặc Unicode không thể lưu trong XLSX.")
    if len(value.encode("utf-16-le")) // 2 <= MAX_CELL_UNITS:
        return [value]
    parts, current, units = [], [], 0
    for char in value:
        size = 2 if ord(char) > 0xFFFF else 1
        if units + size > MAX_CELL_UNITS:
            parts.append("".join(current))
            current, units = [], 0
        current.append(char)
        units += size
    parts.append("".join(current))
    return parts


@dataclass
class Column:
    name: str
    parts: int = 1
    width: int = 12


def fingerprint(path: Path):
    st = path.stat()
    return st.st_dev, st.st_ino, st.st_size, st.st_mtime_ns


def convert(source: Path, output: Path | None = None, *, encoding="utf-8-sig", delimiter=None, force=False) -> dict:
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Thiếu openpyxl. Cài bằng: python -m pip install openpyxl==3.1.5") from exc

    source = Path(source).expanduser().resolve(strict=True)
    if not source.is_file() or source.suffix.lower() not in {".jsonl", ".ndjson", ".csv"}:
        raise ValueError("Input phải là tệp .jsonl, .ndjson hoặc .csv.")
    output = Path(output).expanduser().absolute() if output is not None else Path.cwd() / (source.stem + ".xlsx")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Tệp đầu ra phải có đuôi .xlsx.")
    if output.resolve() == source:
        raise ValueError("Đường dẫn output không được trùng input.")
    if not output.parent.is_dir():
        raise ValueError(f"Thư mục đích không tồn tại: {output.parent}")
    if output.exists() and not force:
        raise FileExistsError(f"Đã có tệp: {output}. Dùng --force để ghi đè hoặc -o để đổi tên.")
    if delimiter is not None and (len(delimiter) != 1 or delimiter in "\r\n\0"):
        raise ValueError("--delimiter phải là một ký tự, hoặc \\t cho tab.")

    initial = fingerprint(source)
    columns, count = {}, 0
    for line, row in records(source, encoding, delimiter):
        count += 1
        for name, value in row.items():
            if name not in columns:
                cell_parts(name, "Tên cột")
                if len(name.encode('utf-16-le')) // 2 > MAX_CELL_UNITS:
                    raise ValueError("Tên cột dài hơn giới hạn Excel.")
                columns[name] = Column(name)
            parts = cell_parts(value, f"Dòng {line}, cột {name}")
            col = columns[name]
            col.parts = max(col.parts, len(parts))
            width = max((len(s) for s in str(parts[0] if parts[0] is not None else "").splitlines()), default=0)
            col.width = min(72, max(col.width, min(len(name) + 2, 40), width + 2))
        if sum(c.parts for c in columns.values()) > MAX_COLUMNS:
            raise ValueError("Dữ liệu sau khi làm phẳng vượt giới hạn 16.384 cột Excel.")
    if not count:
        raise ValueError("Input không có bản ghi dữ liệu.")
    if fingerprint(source) != initial:
        raise ValueError("Input đã thay đổi khi đang đọc. Hãy chạy lại.")

    headers, used = [], set()
    for col in columns.values():
        for part in range(col.parts):
            name = col.name if part == 0 else f"{col.name} [part {part + 1}]"
            headers.append(unique_name(name.replace("\r\n", "\n").replace("\r", "\n"), used))
    book = Workbook(write_only=True)
    book.properties.creator = "BKQuiz"
    font = Font(name="Calibri", size=11)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="17365D")
    align = Alignment(vertical="top", wrap_text=True)

    def make_cell(sheet, value, header=False):
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = "s"  # Preserve source text, including strings beginning with '='.
            cell.number_format = "@"
        cell.font = header_font if header else font
        cell.alignment = align
        if header:
            cell.fill = header_fill
        return cell

    def new_sheet(number):
        sheet = book.create_sheet("Data" if number == 1 else f"Data_{number}")
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 36
        index = 1
        for col in columns.values():
            for _ in range(col.parts):
                sheet.column_dimensions[get_column_letter(index)].width = col.width
                index += 1
        sheet.append([make_cell(sheet, h, True) for h in headers])
        return sheet

    handle, temp_name = tempfile.mkstemp(prefix=".xlsx-export-", suffix=".xlsx", dir=output.parent)
    os.close(handle)
    temp_path = Path(temp_name)
    sheet_count, sheet_rows, written = 1, 0, 0
    sheet = new_sheet(sheet_count)
    try:
        for line, row in records(source, encoding, delimiter):
            if set(row) - set(columns):
                raise ValueError("Input thay đổi cột khi đang xuất. Hãy chạy lại.")
            if sheet_rows == MAX_DATA_ROWS:
                sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet_rows + 1}"
                sheet_count += 1
                sheet, sheet_rows = new_sheet(sheet_count), 0
            values = []
            for col in columns.values():
                parts = cell_parts(row.get(col.name), f"Dòng {line}, cột {col.name}")
                if len(parts) > col.parts:
                    raise ValueError("Input thay đổi độ dài ô khi đang xuất. Hãy chạy lại.")
                values.extend(parts + [None] * (col.parts - len(parts)))
            sheet.append([make_cell(sheet, v) for v in values])
            sheet_rows += 1
            written += 1
        if written != count or fingerprint(source) != initial:
            raise ValueError("Input đã thay đổi khi đang xuất. Hãy chạy lại.")
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet_rows + 1}"
        book.save(temp_path)
        if force:
            os.replace(temp_path, output)
        else:
            # Exclusive creation also protects against a file created during conversion.
            created = False
            try:
                with output.open("xb") as dst:
                    created = True
                    with temp_path.open("rb") as src:
                        shutil.copyfileobj(src, dst)
            except BaseException:
                if created:
                    output.unlink(missing_ok=True)
                raise
    finally:
        for ws in book.worksheets:
            if not ws.closed:
                ws.close()
        book.close()
        temp_path.unlink(missing_ok=True)
    return {"output": str(output), "records": written, "columns": len(headers), "sheets": sheet_count,
            "split_text_columns": sum(c.parts > 1 for c in columns.values())}


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Chuyển JSONL/CSV thành XLSX. Mặc định lưu vào thư mục hiện tại của terminal.")
    parser.add_argument("input", type=Path, help="Đường dẫn tệp .jsonl, .ndjson hoặc .csv")
    parser.add_argument("-o", "--output", type=Path, help="Tên/đường dẫn output; mặc định: ./<tên input>.xlsx")
    parser.add_argument("--encoding", default="utf-8-sig", help="Encoding input (mặc định: utf-8-sig; ví dụ cp1258)")
    parser.add_argument("--delimiter", help="Dấu phân cách CSV; tự nhận diện nếu bỏ trống. Dùng \\t cho tab.")
    parser.add_argument("--force", action="store_true", help="Cho phép ghi đè output đã có")
    args = parser.parse_args(argv)
    try:
        result = convert(args.input, args.output, encoding=args.encoding,
                         delimiter="\t" if args.delimiter == "\\t" else args.delimiter, force=args.force)
    except (OSError, ValueError, RuntimeError, LookupError, RecursionError) as exc:
        print(f"Lỗi: {exc}", file=sys.stderr)
        return 2
    print(f"Đã tạo: {result['output']}")
    print(f"{result['records']} bản ghi | {result['columns']} cột | {result['sheets']} sheet")
    if result['split_text_columns']:
        print(f"Có {result['split_text_columns']} trường dài được chia sang cột [part 2], [part 3]...")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
